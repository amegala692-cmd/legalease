from __future__ import annotations

import base64
import io
import re
import os
import tempfile
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from PIL import Image

from backend.utils.text_utils import sanitize_text, safe_filename


def decode_logo(data_uri: Optional[str]) -> Optional[bytes]:
    if not data_uri:
        return None
    if "," in data_uri:
        _, encoded = data_uri.split(",", 1)
    else:
        encoded = data_uri
    try:
        return base64.b64decode(encoded)
    except Exception:
        return None


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def _style_doc(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)


def format_docx(text: str, doc_type: str, company_name: str = "LegalEase", footer: str = "", effective_date: str = "", terms: Optional[list[str]] = None, logo_data: Optional[bytes] = None) -> bytes:
    doc = Document()
    _style_doc(doc)

    if logo_data:
        try:
            image_stream = io.BytesIO(logo_data)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(image_stream, width=Inches(1.1))
        except Exception:
            pass

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(sanitize_text(doc_type).upper())
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"{company_name} | Effective Date: {effective_date or '[Not specified]'}").italic = True

    lines = sanitize_text(text).splitlines()
    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue
        if re.match(r"^\d+[.)]\s+", stripped) or (stripped.isupper() and len(stripped) < 100 and re.search(r"[A-Z]", stripped)):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(stripped)
            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
        else:
            p = doc.add_paragraph(stripped)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.08

    if terms:
        doc.add_paragraph()
        h = doc.add_paragraph()
        h.add_run("Key Terms Table").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "#"
        table.rows[0].cells[1].text = "Term / Condition"
        for idx, term in enumerate(terms, 1):
            cells = table.add_row().cells
            cells[0].text = str(idx)
            cells[1].text = sanitize_text(term)

    for section in doc.sections:
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_p.text = footer or f"{company_name} | Informational draft — not legal advice"

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalPDF(FPDF):
    def __init__(self, company_name: str, footer: str, logo_data: Optional[bytes]):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.company_name = company_name
        self.footer_text = footer or f"{company_name} | Informational draft — not legal advice"
        self.logo_data = logo_data
        self._logo_path = None

    def header(self):
        self.set_font("Helvetica", "B", 9)
        if self.logo_data:
            try:
                image = Image.open(io.BytesIO(self.logo_data)).convert("RGB")
                with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                    image.save(tmp.name, format="PNG")
                    path = tmp.name
                try:
                    self.image(path, x=95, y=8, w=20)
                finally:
                    try:
                        os.unlink(path)
                    except OSError:
                        pass
                self.set_y(30)
            except Exception:
                self.set_y(12)
        else:
            self.set_y(12)
        self.cell(0, 6, self.company_name, align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.ln(8)

    def footer(self):
        self.set_y(-18)
        self.set_font("Helvetica", "I", 8)
        self.multi_cell(0, 4, _pdf_ascii(self.footer_text), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)


def _pdf_ascii(text: str) -> str:
    return text.encode("latin-1", "replace").decode("latin-1")


def format_pdf(text: str, doc_type: str, company_name: str = "LegalEase", footer: str = "", effective_date: str = "", logo_data: Optional[bytes] = None) -> bytes:
    pdf = LegalPDF(company_name=company_name, footer=footer, logo_data=logo_data)
    pdf.set_auto_page_break(auto=True, margin=22)
    pdf.add_page()
    pdf.set_title(sanitize_text(doc_type))

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 9, _pdf_ascii(sanitize_text(doc_type).upper()), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", "I", 9)
    pdf.multi_cell(0, 6, _pdf_ascii(f"Effective Date: {effective_date or '[Not specified]'}"), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(5)

    for raw_line in sanitize_text(text).splitlines():
        line = raw_line.strip()
        if not line:
            pdf.ln(2)
            continue
        is_heading = bool(re.match(r"^\d+[.)]\s+", line)) or (line.isupper() and len(line) < 100 and re.search(r"[A-Z]", line))
        pdf.set_font("Helvetica", "B" if is_heading else "", 11 if is_heading else 10.5)
        pdf.multi_cell(0, 6.2 if not is_heading else 7, _pdf_ascii(line), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        if is_heading:
            pdf.ln(1)
    result = pdf.output(dest="S")
    return bytes(result)
